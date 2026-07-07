import openpyxl, warnings
from openpyxl.styles import Font, PatternFill
warnings.filterwarnings('ignore')
wb = openpyxl.load_workbook('04_Taakvelden_2027_backup_00_origineel.xlsx')
taakvelden=[n for n in wb.sheetnames if n and n[0]=='T' and any(c.isdigit() for c in n)]
# Voortgang
if 'Voortgang' in wb.sheetnames: del wb['Voortgang']
ws=wb.create_sheet('Voortgang',1)
bold=Font(bold=True); fill=PatternFill('solid',fgColor='D9E1F2')
for c,h in enumerate(['Taakveld','Status','Datum laatste bewerking','Eerstvolgende stap','Aandachtspunten'],1):
    cell=ws.cell(1,c,h); cell.font=bold; cell.fill=fill
for col,w in {'A':10,'B':28,'C':22,'D':38,'E':62}.items(): ws.column_dimensions[col].width=w
r=2
for tv in taakvelden:
    ws.cell(r,1,tv)
    if tv=='T01':
        ws.cell(r,2,'bezig'); ws.cell(r,4,'Fase 1 stap 2-3: uniformeren + herbereken-controle'); ws.cell(r,5,'Pilot / standaardsjabloon.')
    elif tv=='T82':
        ws.cell(r,2,'nog te doen (black box, Bijlage B.0)'); ws.cell(r,4,'Interface-totalen verifieren na akkoord collega'); ws.cell(r,5,'Niet uniformeren; 3 totaalregels/jaar + uitgelijnde kolommen. Aandeel 41100/47501 handmatige aansluiting.')
    elif tv in ('T04L','T04B','T04W'):
        ws.cell(r,2,'nog te doen (complex, uitgesteld)'); ws.cell(r,4,'Later: kostenverdeelstaat/overhead, samenhang Fase 3'); ws.cell(r,5,'Complexe kostenverdeelstaat overhead. Bewaren voor later.')
    else:
        ws.cell(r,2,'nog te doen'); ws.cell(r,4,'Wacht op goedkeuring T01')
    r+=1
nr=len(taakvelden)+4
ws.cell(nr,1,'BESLISPUNTEN (voorlopige keuzes, te bevestigen bij controle T01):').font=bold
for i,k in enumerate(['1. Historie 2025/2026: BEVROREN als vaste waarden (anker).','2. Controleritme na T01: batches van ~5.','3. Inwonersreeks: nog aan te leveren (2021-2030); inwonerposten T01 voorlopig eigen onderbouwing, later koppelen.','4. Inlees CODA: bestaand sjabloon + 2030 (Fase 4).','5. Overhead fte: input nog aanleveren (Fase 3).','6. Opruimen Blad1/2/3 + dubbele inhoudsopgave + oude draaitabellen: na afronding.','7. T82: black box (B.0); volledige uniformering pas na akkoord collega.','8. Bladbeveiliging: nog te bepalen.','9. Volgorde koppelen staten: nog te bepalen.','10. Kolomgroei jaarovergang: nog te bepalen.']):
    ws.cell(nr+1+i,1,k)
nr2=nr+12
ws.cell(nr2,1,'OPENSTAANDE AANSLUITCONTROLES STATEN (Fase 2, handmatig tot koppeling):').font=bold
for i,t in enumerate(['41100 salarissen','471/871 reserves en voorzieningen','473/873 activa','475/875 verrekening']):
    ws.cell(nr2+1+i,1,'- totaal '+t+' over alle taakvelden vs staat: nog te leggen')
# Indexering annotatie
wi=wb['Indexering']
wi['A3']='INDEXREEKSEN (bron; 2021-2030 in D:M)'; wi['A3'].font=bold
if wi['B6'].value in (None,''): wi['B6']='Gewogen index'
wi['A23']='AANTAL INWONERS (nog aan te leveren, beslispunt 3; 2021-2030 in D:M, met bronvermelding)'; wi['A23'].font=bold
wi['B24']='Aantal inwoners'
for c,jr in zip(range(4,14),range(2021,2031)): wi.cell(23,c,jr)
wi['A26']='Toelichting: "index conform kaderbrief" (rij 16-20) later samenvoegen/koppelen met bovenstaande reeksen.'
wb.save('04_Taakvelden_2027_na_verwerken_huren.xlsx')
print('Werkbestand schoon herbouwd vanaf backup + Fase 0 toegepast. Sheets:',len(wb.sheetnames))
