# Content-only transform op wegwerp-kopie -> levert wijzigingslijst voor XML-chirurgie.
import openpyxl, json, re, warnings, shutil
warnings.filterwarnings('ignore')
snap=json.load(open('snapshot_fase0_waarden.json'))['T01']
shutil.copy('04_Taakvelden_2027_backup_00_origineel.xlsx','/tmp/temp_content.xlsx')
wb=openpyxl.load_workbook('/tmp/temp_content.xlsx'); ws=wb['T01']

# hulpblok labels + sociaal-domein rij
ws['P1']='reeks'; ws['P2']='Prijzen'; ws['P3']='Lonen'; ws['P4']='Gewogen'; ws['P5']='Sociaal domein'
for h,i in zip('QRSTUVW','GHIJKLM'): ws[f'{h}5']=f'=Indexering!{i}7'

# STAP A: labels->tekst, A/B->waarde, stijl =+ -> =
for row in ws.iter_rows():
    for cell in row:
        v=cell.value
        if not (isinstance(v,str) and v.startswith('=')): continue
        cached=snap.get(cell.coordinate); col=cell.column_letter
        if isinstance(cached,str): cell.value=cached; continue
        if col in ('A','B'): cell.value=cached; continue
        if v.startswith('=+'): cell.value='='+v[2:]

# STAP B/C: onderlaag index-formules -> A.4 met H
NAME={2:'Prijzen',3:'Lonen',4:'Gewogen',5:'Sociaal domein'}; HULP={'L':'T','M':'U','N':'V','O':'W'}
for r in range(219, ws.max_row+1):
    fL=ws.cell(r,12).value
    if not (isinstance(fL,str) and fL.startswith('=')): continue
    m=re.search(r'\*\s*\$?([A-Z]+)\$?(\d)', fL)
    if not m: continue
    letter,digit=m.group(1),int(m.group(2))
    if letter!='T' or digit not in NAME: continue
    hc=ws.cell(r,8)
    if hc.value not in (None,''): continue   # H bezet -> laat origineel
    hc.value=NAME[digit]
    for c,dcol in ((12,'L'),(13,'M'),(14,'N'),(15,'O')):
        prev=chr(ord(dcol)-1)+str(r); hy=HULP[dcol]
        ws.cell(r,c).value=(f'=IF($H{r}="Geen",{prev},ROUND({prev}*INDEX({hy}$2:{hy}$5,MATCH($H{r},$P$2:$P$5,0)),0))')

# STAP D: staat-toelichting salaris
ws.cell(510,5).value='zie staat van salarissen (nu loongeindexeerd in blad; later koppeling)'
ws.cell(519,5).value='zie staat van salarissen (nu loongeindexeerd in blad; later koppeling)'
ws.cell(510,8).value='Lonen'; ws.cell(519,8).value='Lonen'
# STAP F: samengestelde invoer r229 uitsplitsen (toelichting) + vaste waarde
ws.cell(229,10).value=snap.get('J229')
ws.cell(230,5).value='(uitsplitsing jaarfactuur Notubiz 2025: 43.940 + 15.751 = 59.691)'
# STAP G: spec-header onderlaag
ws.cell(220,5).value=('Specificatieplicht: categorieen 42/43/44/45 en 82/83/84/85 via detailregels. '
 'Standaardtoelichtingen: 41100 zie staat van salarissen - 471/871 reserves/voorzieningen - '
 '473/873 activa - 475/875 verrekening.')
# meldingsregel
ws.cell(8,4).value=('MELDING T01: geuniformeerd volgens Bijlage A (3 lagen, indexkeuze kolom H, A.4-patroon). '
 '1 afwijking >EUR1 gemeld: Jeugdraad, afrondingsverschil tot EUR1,39 in 2030 (uniform AFRONDEN; was ongeafrond). '
 'Totaal lasten 2027 -EUR1 (binnen tolerantie). Waarden verder gelijk aan snapshot.')
# H/I kolomkoppen waar vrij (niet-merged)
from openpyxl.cell.cell import MergedCell
for hr in (32,62,76):
    for cc,lbl in ((8,'Indexkeuze'),(9,'Btw/BCF')):
        t=ws.cell(hr,cc)
        if not isinstance(t,MergedCell): t.value=lbl

wb.save('/tmp/temp_content.xlsx')
print('content-transform klaar')
