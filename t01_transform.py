import openpyxl, json, re, warnings
from openpyxl.styles import Font, PatternFill
warnings.filterwarnings('ignore')
snap = json.load(open('snapshot_fase0_waarden.json'))['T01']
def sv(coord): return snap.get(coord)
wb = openpyxl.load_workbook('04_Taakvelden_2027_na_verwerken_huren.xlsx')
ws = wb['T01']
YEL = PatternFill('solid', fgColor='FFF2CC')   # geel = handmatige correctie
BLU = Font(color='1F4E78')                       # blauw = invoer (indicatief)
log = {'labels':0,'ab':0,'style':0,'idx':0,'H':0,'Hbezet':0,'toel':0,'corrmark':0}

# ---------- STAP A: labels->tekst, A/B->waarde, stijl =+ -> = ----------
for row in ws.iter_rows():
    for cell in row:
        v=cell.value
        if not (isinstance(v,str) and v.startswith('=')): continue
        cached=sv(cell.coordinate)
        col=cell.column_letter
        if isinstance(cached,str):
            cell.value=cached; log['labels']+=1; continue
        if col in ('A','B'):
            cell.value=cached; log['ab']+=1; continue
        if v.startswith('=+'):
            cell.value='='+v[2:]; log['style']+=1

# ---------- STAP B/C: onderlaag index-formules -> A.4 met H ----------
NAME={2:'Prijzen',3:'Lonen',4:'Gewogen',5:'Sociaal domein'}
HULP={'L':'T','M':'U','N':'V','O':'W'}   # datakolom -> hulpblok-jaarkolom (zelfde jaar)
def hcol_ok(letter): return letter in 'QRSTUVW'
onder_start=219
for r in range(onder_start, ws.max_row+1):
    # bepaal indexsoort uit ORIGINELE L-formule (via snapshot? nee: uit huidige formule)
    fL = ws.cell(r,12).value  # kolom L
    if not (isinstance(fL,str) and fL.startswith('=')): continue
    # zoek vermenigvuldiging met hulpblok-jaarcel: *<letter><digit> met letter=T (voor L) digit 2-5
    m=re.search(r'\*\s*\$?([A-Z]+)\$?(\d)', fL)
    if not m: continue
    letter, digit = m.group(1), int(m.group(2))
    if letter!='T' or digit not in NAME: continue   # alleen echte index-formules op jaarcel T (2027)
    Hname=NAME[digit]
    # H zetten indien vrij
    hc=ws.cell(r,8)  # kolom H
    if hc.value not in (None,''):
        log['Hbezet']+=1; continue   # H bezet door sub-berekening -> niet herschrijven (layout-item)
    hc.value=Hname; log['H']+=1
    for c,dcol in ((12,'L'),(13,'M'),(14,'N'),(15,'O')):
        prev=chr(ord(dcol)-1)+str(r)
        hy=HULP[dcol]
        ws.cell(r,c).value=(f'=IF($H{r}="Geen",{prev},'
                            f'ROUND({prev}*INDEX({hy}$2:{hy}$5,MATCH($H{r},$P$2:$P$5,0)),0))')
        log['idx']+=1

# ---------- STAP D: staat-toelichtingen (salaris 41100) ----------
# onderlaag 41100-blokken: r504 (Directe toerekening bestuur), r513 (regionale)
for tr in (509, 518):  # onder de 41100-totaalregels
    if ws.cell(tr,5).value and 'zie staat' not in str(ws.cell(tr,5).value):
        pass
ws.cell(510,5).value='zie staat van salarissen (nu loongeindexeerd in blad; later koppeling staat van salarissen)'
ws.cell(519,5).value='zie staat van salarissen (nu loongeindexeerd in blad; later koppeling staat van salarissen)'
ws.cell(510,8).value='Lonen'; ws.cell(519,8).value='Lonen'
log['toel']+=2

# ---------- STAP E: harde correctieregels markeren + H=Geen ----------
# W02-2025 (r238), A03-2025 (r239), A03-2025 werkbudget (r295), W02/2025 VNG (r359),
# W01/23 (r285), Best-Oirschot W-regels (r493-498), staffunctie/schaalsprong (r517,521,522)
corr_rows=[238,239,285,295,359,493,494,495,496,497,498,517,521,522]
for r in corr_rows:
    for c in range(10,16):  # J..O
        ws.cell(r,c).fill=YEL
    if ws.cell(r,8).value in (None,''):
        ws.cell(r,8).value='Geen'
    log['corrmark']+=1

# ---------- STAP F: samengestelde invoer r229 (=43940+15751) ----------
# J229 was =43940+15751 (2025, bevroren). Zet vaste waarde + uitsplitsing in toelichting.
ws.cell(229,10).value=snap.get('J229')  # 59691
ws.cell(230,5).value='(uitsplitsing jaarfactuur Notubiz 2025: 43.940 + 15.751 = 59.691)'

# ---------- STAP G: standaard specificatie-/toelichtingsregel bovenaan onderlaag ----------
ws.cell(220,5).value=('Specificatieplicht: categorieen 42/43/44/45 en 82/83/84/85 via detailregels. '
 'Standaardtoelichtingen: 41100 zie staat van salarissen - 471/871 zie staat van reserves en '
 'voorzieningen - 473/873 zie staat van activa - 475/875 verrekening.')
ws.cell(220,5).font=Font(italic=True)

# ---------- H/I kolomkoppen ----------
from openpyxl.cell.cell import MergedCell
for hdr_row in (32,62,76):
    for cc,lbl in ((8,'Indexkeuze'),(9,'Btw/BCF')):
        tgt=ws.cell(hdr_row,cc)
        if isinstance(tgt,MergedCell):  # sla samengevoegde over
            continue
        tgt.value=lbl

wb.save('04_Taakvelden_2027_na_verwerken_huren.xlsx')
print(json.dumps(log))
